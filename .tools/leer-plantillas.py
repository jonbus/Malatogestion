"""Saca las tres plantillas de turnos, enteras."""
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

RUTA = r"C:\Users\Usuario\Desktop\APP Malato Documentación\PlantillasHorarios.xlsx"
wb = load_workbook(RUTA, data_only=True)

for nombre in wb.sheetnames:
    ws = wb[nombre]
    if ws.max_row <= 1:
        continue
    print(f"\n{'='*70}")
    print(f"  {nombre}   ({ws.max_row} filas x {ws.max_column} columnas)")
    print("=" * 70)
    for r in range(1, ws.max_row + 1):
        celdas = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                celdas.append("")
            elif isinstance(v, float):
                celdas.append(f"{v:g}")
            else:
                celdas.append(str(v).strip())
        if any(celdas):
            print(f"  r{r:<3} | " + " | ".join(celdas))
