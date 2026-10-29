"""Verificacion final de la copia anonimizada y recalculada contra el original.

Comprueba tres cosas:
  A) no queda ningun nombre propio
  B) las FORMULAS son identicas celda a celda
  C) los VALORES CALCULADOS coinciden donde el original los tenia
"""
import re
from openpyxl import load_workbook

O = r"C:\Users\Usuario\Desktop\APP Malato Documentación\EjemploHorariosVacacionesVerano2026.xlsx"
C = r"C:\Users\Usuario\Desktop\APP Malato\docs\referencia\horarios-verano-2026-anonimizado.xlsx"

fo, fc = load_workbook(O, data_only=False), load_workbook(C, data_only=False)
vo, vc = load_workbook(O, data_only=True), load_workbook(C, data_only=True)

# Nombres a buscar: se derivan del original por posicion, sin imprimirlos.
ws = fo["VacasTodas2026"]
nombres = [ws.cell(row=r, column=2).value for r in (438, 440, 442, 447, 449)]
nombres = [n.strip() for n in nombres if isinstance(n, str)]

# A) restos
restos = 0
for hoja in fc.sheetnames:
    for fila in fc[hoja].iter_rows():
        for c in fila:
            if isinstance(c.value, str) and any(n in c.value for n in nombres):
                restos += 1

# B y C) recorrido comparativo
form_distintas = 0
valor_distinto = 0
valor_antes_si_ahora_no = 0
comparadas = 0
for hoja in fo.sheetnames:
    so, sc = fo[hoja], fc[hoja]
    so_v, sc_v = vo[hoja], vc[hoja]
    for r in range(1, max(so.max_row, sc.max_row) + 1):
        for col in range(1, max(so.max_column, sc.max_column) + 1):
            a, b = so.cell(row=r, column=col).value, sc.cell(row=r, column=col).value
            comparadas += 1
            if isinstance(a, str) and any(n in a for n in nombres):
                continue                      # cambio autorizado: etiqueta
            if a != b:
                form_distintas += 1
                if form_distintas <= 4:
                    print(f"  formula distinta en {hoja}!{sc.cell(row=r, column=col).coordinate}: {a!r} -> {b!r}")
            av, bv = so_v.cell(row=r, column=col).value, sc_v.cell(row=r, column=col).value
            if isinstance(av, (int, float)):
                if not isinstance(bv, (int, float)):
                    valor_antes_si_ahora_no += 1
                    if valor_antes_si_ahora_no <= 4:
                        print(f"  valor perdido en {hoja}!{sc.cell(row=r, column=col).coordinate}: {av!r} -> {bv!r}")
                elif abs(av - bv) > 1e-9:
                    valor_distinto += 1
                    if valor_distinto <= 4:
                        print(f"  valor distinto en {hoja}!{sc.cell(row=r, column=col).coordinate}: {av!r} -> {bv!r}")

print(f"A · nombres que quedan: {restos}")
print(f"B · celdas comparadas: {comparadas} · formulas distintas: {form_distintas}")
print(f"C · valores numericos perdidos: {valor_antes_si_ahora_no} · valores distintos: {valor_distinto}")

w = vc["VacasTodas2026"]
print("\nMuestra de resultados calculados en la copia:")
print(f"  R4={w.cell(row=4, column=18).value}  R12={w.cell(row=12, column=18).value}  "
      f"R438={w.cell(row=438, column=18).value}")
print(f"  H.Trabajadas: F635={w.cell(row=635, column=6).value}  "
      f"F636={w.cell(row=636, column=6).value}  F637={w.cell(row=637, column=6).value}")

ok = restos == 0 and form_distintas == 0 and valor_distinto == 0 and valor_antes_si_ahora_no == 0
print("\nRESULTADO:", "OK — copia fiel al original salvo las etiquetas" if ok else "REVISAR")
