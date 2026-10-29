"""Comprueba si los valores calculados de las formulas sobrevivieron a la copia."""
from openpyxl import load_workbook

O = r"C:\Users\Usuario\Desktop\APP Malato Documentación\EjemploHorariosVacacionesVerano2026.xlsx"
D = r"C:\Users\Usuario\Desktop\APP Malato\docs\referencia\horarios-verano-2026-anonimizado.xlsx"

for etq, ruta in (("ORIGINAL", O), ("COPIA   ", D)):
    w = load_workbook(ruta, data_only=True)["VacasTodas2026"]
    filas = [4, 12, 438]          # columnas R (18) = totales semanales
    tot = [635, 636, 637]         # columna F (6) = horas trabajadas
    sr = "  ".join(f"R{r}={w.cell(row=r, column=18).value!r}" for r in filas)
    sf = "  ".join(f"F{r}={w.cell(row=r, column=6).value!r}" for r in tot)
    print(f"{etq}  {sr}")
    print(f"{'':9}  {sf}")
