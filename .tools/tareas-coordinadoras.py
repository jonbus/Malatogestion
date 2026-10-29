"""Tareas de coordinadoras: lista distinta, dias, y solapes con personas de tienda."""
import sys
from collections import defaultdict
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook(r".tools\tareas.xlsx", data_only=True)


def parsear(hoja):
    ws = wb[hoja]
    dia = None
    salida = defaultdict(list)
    for r in range(1, ws.max_row + 1):
        v1 = ws.cell(row=r, column=1).value
        v2 = ws.cell(row=r, column=2).value
        v3 = ws.cell(row=r, column=3).value
        cabecera = None
        for v in (v1, v2):
            if isinstance(v, str) and "ORDINARIAS" in v.upper():
                cabecera = v.strip().upper().replace("TAREAS ORDINARIAS", "").strip()
        if cabecera:
            dia = cabecera
            continue
        if not isinstance(v2, str) or not v2.strip():
            continue
        t = v2.strip()
        if t.upper() in ("TAREAS", "\u221a") or t.startswith("*") or t.upper() == "FICHAR":
            continue
        salida[dia].append((t, str(v3).strip() if v3 else ""))
    return salida


coord = parsear("Tareas Coordinadoras")
info = defaultdict(lambda: {"dias": set(), "obs": ""})
for dia, lista in coord.items():
    if not dia:
        continue
    for t, obs in lista:
        info[t]["dias"].add(dia[:3])
        if obs and not info[t]["obs"]:
            info[t]["obs"] = obs

print("=" * 78)
print(f"COORDINADORAS · {len(info)} tareas distintas")
print("=" * 78)
for t in sorted(info, key=lambda x: (-len(info[x]["dias"]), x.lower())):
    d = sorted(info[t]["dias"])
    obs = f"   [{info[t]['obs'][:40]}]" if info[t]["obs"] else ""
    print(f"  {len(d)}d {'/'.join(d):<22} | {t}{obs}")

# --- solapes con personas de tienda
op = parsear("1TareasDiariasMañanaTardeOperar")
op2 = parsear("2TareasDiariasMañanaTardeOperar")
tops = {t for v in op.values() for t, _ in v} | {t for v in op2.values() for t, _ in v}
tco = set(info)


def normalizar(t):
    t = t.lower()
    for ch in "áéíóúü(),./…*-":
        t = t.replace(ch, " ")
    return " ".join(t.split())


nop = {}
for t in tops:
    nop.setdefault(normalizar(t), []).append(t)

print("\n" + "=" * 78)
print("SOLAPES CON TIENDAS (misma tarea en los dos puestos)")
print("=" * 78)
for t in sorted(tco, key=str.lower):
    n = normalizar(t)
    if n in nop:
        print(f"  {t}")
        for o in nop[n]:
            print(f"      = personas de tienda: {o}")

print("\n--- tareas parecidas por palabras clave ---")
CLAVES = ["fruta", "fregona", "monta", "leche", "yogur", "lineal", "recepci", "tickar",
          "boller", "pan", "pollo", "vending", "danone", "cesto", "trastienda", "cámara"]
for clave in CLAVES:
    c = [t for t in tco if clave in t.lower()]
    o = [t for t in tops if clave in t.lower()]
    if c and o:
        print(f"\n  [{clave}]")
        for x in c:
            print(f"     coord:     {x}")
        for x in o:
            print(f"     personas de tienda: {x}")
