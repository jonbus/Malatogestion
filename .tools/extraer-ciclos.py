"""Extrae los ciclos de turnos DEDUCIDOS de las plantillas reales, tal cual estan.

Salida: para cada persona, que turno tiene cada dia de la semana, en cada semana
del ciclo. Es lo que hay que validar con la propiedad.
"""
import re
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
R = r"C:\Users\Usuario\Desktop\APP Malato Documentación\PlantillasHorarios.xlsx"
DIAS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"]
ABREV = ["L", "M", "X", "J", "V", "S", "D"]
wb = load_workbook(R, data_only=True)

# (hoja, filas de personas por semana) -> el bloque de cada persona son 2 filas
HOJAS = [
    ("Turnos Tienda", ["OP 1", "OP 2", "OP 3"]),
    ("Turnos EncargadasIndividualSinC", ["COORD 1", "COORD2"]),
    ("Turnos Gerencia", ["GER"]),
]

datos = {}          # (persona) -> {semana: {dia: [turnos]}}
for hoja, personas in HOJAS:
    ws = wb[hoja]
    semana = None
    persona = None
    bloque_duplicado = False
    vistos = set()
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=8):
        et = fila[0].value
        b1 = str(fila[1].value or "").strip().upper()
        if b1 in DIAS:                       # cabecera de semana
            if et is not None:
                m = re.search(r"\d", str(et))
                if m:
                    semana = int(m.group())
            persona = None
            bloque_duplicado = False
            continue
        if et is not None:
            e = str(et).strip().upper()
            if re.match(r"^\d", e) or "SEMANA" in e:
                m = re.search(r"\d", e)
                if m:
                    semana = int(m.group())
                persona = None
                continue
            if e in [p.upper() for p in personas]:
                persona = e
                # Cada persona ocupa DOS filas: la etiqueta esta en la primera y el
                # turno puede estar repartido entre las dos. Hay que leer el bloque
                # ENTERO. Lo que se salta son los BLOQUES repetidos, no las filas.
                # OJO: aqui NO se hace continue, porque la etiqueta y los turnos
                # comparten la misma fila.
                clave = (persona, semana)
                bloque_duplicado = clave in vistos
                vistos.add(clave)
            else:
                persona = None
                continue
        if persona is None or semana is None or bloque_duplicado:
            continue
        for idx, c in enumerate(fila[1:8]):
            t = str(c.value or "").strip()
            if not t or t.upper() in ("VAC", "LIBRE", "$$$"):
                continue
            datos.setdefault(persona, {}).setdefault(semana, {}).setdefault(idx, []).append(t)

print("=" * 78)
print("CICLOS DEDUCIDOS DE TUS PLANTILLAS (esto es lo que hay que validar)")
print("=" * 78)
for persona in ["OP 1", "OP 2", "OP 3", "COORD 1", "COORD2", "GER"]:
    if persona not in datos:
        continue
    print(f"\n### {persona}")
    for sem in sorted(datos[persona]):
        fila = []
        for idx in range(7):
            turnos = datos[persona][sem].get(idx, [])
            fila.append(f"{ABREV[idx]}={' + '.join(turnos) if turnos else '--'}")
        print(f"   semana {sem}: " + "  ".join(fila))
