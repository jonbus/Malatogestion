"""Lee la clasificación de fruta y verdura que ha completado la propiedad."""
import sys
from collections import Counter

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook("docs/fruta-verdura-familias.xlsx", data_only=True)
ws = wb["Clasificar"]

filas = []
for r in range(5, ws.max_row + 1):
    ci = ws.cell(row=r, column=1).value
    prod = ws.cell(row=r, column=2).value
    fam = ws.cell(row=r, column=3).value
    nota = ws.cell(row=r, column=4).value
    if prod:
        filas.append({
            "ci": "" if ci is None else str(ci).strip(),
            "prod": str(prod).strip(),
            "fam": "" if fam is None else str(fam).strip(),
            "nota": "" if nota is None else str(nota).strip(),
        })

print(f"  filas: {len(filas)}")
c = Counter(f["fam"] for f in filas)
print("\n  === FAMILIAS ===")
for k, v in c.most_common():
    etiqueta = k if k else "(SIN FAMILIA)"
    print(f"    {etiqueta:<24} {v}")

sin_fam = [f for f in filas if not f["fam"]]
print(f"\n  === SIN FAMILIA ({len(sin_fam)}) ===")
for f in sin_fam:
    print(f"    [{f['ci']:<10}] {f['prod']}")

sin_ci = [f for f in filas if f["ci"] in ("", "¿?", "None", "?" )]
print(f"\n  === SIN CI ({len(sin_ci)}) ===")
for f in sin_ci:
    print(f"    [{f['ci']:<10}] {f['prod']:<30} {f['fam']}")

con_nota = [f for f in filas if f["nota"]]
print(f"\n  === CON NOTA ({len(con_nota)}) ===")
for f in con_nota:
    print(f"    {f['prod']:<32} {f['fam']:<22} | {f['nota']}")

# duplicados de CI
cis = [f["ci"] for f in filas if f["ci"] not in ("", "¿?", "None", "?")]
dup = {k: v for k, v in Counter(cis).items() if v > 1}
print(f"\n  === CI REPETIDOS ===")
print("    " + (", ".join(f"{k} (x{v})" for k, v in dup.items()) if dup else "ninguno"))
