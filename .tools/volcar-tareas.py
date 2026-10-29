"""Volcado compacto de una hoja: fila a fila, solo celdas con contenido."""
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
HOJA = sys.argv[1]
DESDE = int(sys.argv[2]) if len(sys.argv) > 2 else 1
HASTA = int(sys.argv[3]) if len(sys.argv) > 3 else 10 ** 6

wb = load_workbook(r".tools\tareas.xlsx", data_only=True)
ws = wb[HOJA]
print(f"=== HOJA {HOJA!r} · filas {DESDE} a {min(HASTA, ws.max_row)} de {ws.max_row} ===")
vacias = 0
for r in range(DESDE, min(HASTA, ws.max_row) + 1):
    celdas = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if v is None or str(v).strip() == "":
            continue
        celdas.append(str(v).replace("\n", " / ").strip())
    if not celdas:
        vacias += 1
        continue
    print(f"  r{r:>3} | " + " || ".join(celdas))
if vacias:
    print(f"  ({vacias} filas vacias omitidas)")
