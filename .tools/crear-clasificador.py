"""Crea el Excel para que la PROPIEDAD clasifique las tareas.

No clasifica nada: deja las columnas vacias con desplegables.
Solo rellena los HECHOS (puesto, tarea, dias que toca, cuantos dias).
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding="utf-8")

cat = json.load(open("docs/tareas-catalogo.json", encoding="utf-8"))

DIAS = {"LUN": "lun", "MAR": "mar", "MIE": "mié", "JUE": "jue",
        "VIE": "vie", "SAB": "sáb", "DOM": "dom"}

TIPOS = [
    ("Rutina", "Se repite todos los días y hay que hacerla"),
    ("Exclusiva del día", "Solo toca ese día concreto, por sus condicionantes"),
    ("Urgente", "Surge, y hay que hacerla ya"),
]

ZONAS = [
    "Lineales", "Yogures", "Fruta", "Verdura", "Congelado", "Bebidas",
    "Cámaras frigoríficas", "Congeladores", "Cámara fruta", "Cámara pollos",
    "Charcutería", "Carne", "Pan y bollería", "Txutxes y paquetería",
    "Caja", "Cocina y asador", "Almacén y trastienda", "Baño",
    "Suelos y exteriores", "Toda la tienda", "Varias",
]

VERDE = "2EC258"
VERDE_OSC = "1E9A45"
SUAVE = "EFFAF2"
CABECERA = PatternFill("solid", fgColor=VERDE_OSC)
TITULO = PatternFill("solid", fgColor=VERDE)
BORDE = Border(*[Side(style="thin", color="D6EDDD")] * 4)

wb = Workbook()

# ---------------------------------------------------------------- hoja 1
ws = wb.active
ws.title = "Clasificar"

ws["A1"] = "CLASIFICAR LAS TAREAS"
ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
ws["A1"].fill = TITULO
ws.merge_cells("A1:G1")
ws.row_dimensions[1].height = 30

ws["A2"] = ("Rellena solo las columnas amarillas: TIPO, ZONA y NOTAS. "
            "Lo demás ya está puesto porque son los hechos de tus plantillas. "
            "No hay nada decidido por mí.")
ws["A2"].font = Font(size=10, italic=True, color="5E6E63")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
ws.merge_cells("A2:G2")
ws.row_dimensions[2].height = 30

CABECERAS = ["Puesto", "Tarea", "Días que toca", "Cuántos días",
             "TIPO", "ZONA", "NOTAS"]
for i, c in enumerate(CABECERAS, 1):
    cel = ws.cell(row=4, column=i, value=c)
    cel.font = Font(bold=True, color="FFFFFF", size=11)
    cel.fill = CABECERA
    cel.alignment = Alignment(horizontal="center", vertical="center")
    cel.border = BORDE
ws.row_dimensions[4].height = 26

AMARILLO = PatternFill("solid", fgColor="FFF8E6")
for i in (5, 6, 7):
    ws.cell(row=4, column=i).fill = PatternFill("solid", fgColor="D97706")

filas = []
for p in cat["puestos"]:
    for t in p["tareas"]:
        d = t.get("dias", [])
        filas.append((p["etiqueta"], t["nombre"],
                      " · ".join(DIAS.get(x, x) for x in d) if d else "según necesidad",
                      len(d), t.get("nota", "")))

# Orden: primero las que tocan mas dias (las rutinas se ven enseguida)
filas.sort(key=lambda f: (f[0] != "Tienda", -f[3], f[1].lower()))

r = 5
for puesto, nombre, dias, nd, nota in filas:
    ws.cell(row=r, column=1, value=puesto).border = BORDE
    c = ws.cell(row=r, column=2, value=nombre)
    c.border = BORDE
    c.alignment = Alignment(wrap_text=False, vertical="center")
    c2 = ws.cell(row=r, column=3, value=dias)
    c2.border = BORDE
    c2.font = Font(size=9, color="5E6E63")
    c3 = ws.cell(row=r, column=4, value=nd if nd else "")
    c3.border = BORDE
    c3.alignment = Alignment(horizontal="center")
    for col in (5, 6, 7):
        cc = ws.cell(row=r, column=col)
        cc.fill = AMARILLO
        cc.border = BORDE
    if nota:
        ws.cell(row=r, column=7).value = None
    r += 1
ULTIMA = r - 1

# ---------------------------------------------------------------- hoja de listas
wl = wb.create_sheet("Listas")
wl["A1"] = "Tipo"
wl["B1"] = "Zona"
for i, (t, _) in enumerate(TIPOS, 2):
    wl.cell(row=i, column=1, value=t)
for i, z in enumerate(ZONAS, 2):
    wl.cell(row=i, column=2, value=z)
wl.sheet_state = "hidden"

# ---------------------------------------------------------------- desplegables
dv_tipo = DataValidation(type="list", formula1="=Listas!$A$2:$A$4",
                         allow_blank=True, showDropDown=False)
dv_tipo.error = "Elige: Rutina, Exclusiva del día o Urgente"
dv_tipo.errorTitle = "Tipo no válido"
ws.add_data_validation(dv_tipo)
dv_tipo.add(f"E5:E{ULTIMA}")

dv_zona = DataValidation(type="list",
                         formula1=f"=Listas!$B$2:$B${len(ZONAS) + 1}",
                         allow_blank=True, showDropDown=False)
ws.add_data_validation(dv_zona)
dv_zona.add(f"F5:F{ULTIMA}")

# ---------------------------------------------------------------- colores por tipo
ws.conditional_formatting.add(f"E5:E{ULTIMA}",
    CellIsRule(operator="equal", formula=['"Rutina"'],
               fill=PatternFill("solid", fgColor="D6EDDD"), font=Font(bold=True, color=VERDE_OSC)))
ws.conditional_formatting.add(f"E5:E{ULTIMA}",
    CellIsRule(operator="equal", formula=['"Exclusiva del día"'],
               fill=PatternFill("solid", fgColor="FDE7C8"), font=Font(bold=True, color="8A4B00")))
ws.conditional_formatting.add(f"E5:E{ULTIMA}",
    CellIsRule(operator="equal", formula=['"Urgente"'],
               fill=PatternFill("solid", fgColor="F7C9C9"), font=Font(bold=True, color="B3261E")))

ws.column_dimensions["A"].width = 15
ws.column_dimensions["B"].width = 62
ws.column_dimensions["C"].width = 26
ws.column_dimensions["D"].width = 9
ws.column_dimensions["E"].width = 19
ws.column_dimensions["F"].width = 22
ws.column_dimensions["G"].width = 34
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:G{ULTIMA}"

# ---------------------------------------------------------------- hoja: que es cada cosa
we = wb.create_sheet("Qué es cada cosa", 1)
we["A1"] = "LOS TRES TIPOS"
we["A1"].font = Font(bold=True, size=16, color="FFFFFF")
we["A1"].fill = TITULO
we.merge_cells("A1:C1")
we.row_dimensions[1].height = 30

we["A3"] = "Tipo"
we["B3"] = "Qué significa"
we["C3"] = "Ejemplo de cómo se reconoce"
for c in ("A3", "B3", "C3"):
    we[c].font = Font(bold=True, color="FFFFFF")
    we[c].fill = CABECERA
    we[c].border = BORDE

EJEMPLOS = {
    "Rutina": "Se hace todos los días. Si un día no se hace, no pasa nada grave… "
              "hasta que pasan cinco y se nota.",
    "Exclusiva del día": "Solo tiene sentido ese día. Hacerla otro día es hacerla mal o no tocaba.",
    "Urgente": "No estaba prevista. Aparece y hay que resolverla, desplazando lo demás.",
}
for i, (t, signif) in enumerate(TIPOS, 4):
    we.cell(row=i, column=1, value=t).font = Font(bold=True)
    we.cell(row=i, column=1).border = BORDE
    we.cell(row=i, column=2, value=signif).alignment = Alignment(wrap_text=True, vertical="top")
    we.cell(row=i, column=2).border = BORDE
    we.cell(row=i, column=3, value=EJEMPLOS[t]).alignment = Alignment(wrap_text=True, vertical="top")
    we.cell(row=i, column=3).border = BORDE
    we.row_dimensions[i].height = 46

we["A8"] = ("OJO CON LA RUTINA: precisamente por ser rutina, se olvida. "
            "Las tareas marcadas como Rutina son las que la aplicación tendrá que poner "
            "DELANTE todos los días, sin que nadie tenga que acordarse de buscarlas.")
we["A8"].font = Font(size=11, italic=True, color="8A4B00")
we["A8"].fill = PatternFill("solid", fgColor="FFF8E6")
we["A8"].alignment = Alignment(wrap_text=True, vertical="center")
we.merge_cells("A8:C8")
we.row_dimensions[8].height = 46

we["A10"] = ("No hace falta rellenar las tres columnas de golpe. "
             "Se puede empezar por el TIPO, que es lo que más cambia la aplicación, "
             "y dejar ZONA y NOTAS para después.")
we["A10"].font = Font(size=10, color="5E6E63")
we["A10"].alignment = Alignment(wrap_text=True, vertical="center")
we.merge_cells("A10:C10")
we.row_dimensions[10].height = 30

we.column_dimensions["A"].width = 20
we.column_dimensions["B"].width = 58
we.column_dimensions["C"].width = 58

# ---------------------------------------------------------------- hoja: resumen
wr = wb.create_sheet("Resumen")
wr["A1"] = "CÓMO VA LA CLASIFICACIÓN"
wr["A1"].font = Font(bold=True, size=16, color="FFFFFF")
wr["A1"].fill = TITULO
wr.merge_cells("A1:D1")
wr.row_dimensions[1].height = 30

wr["A3"] = "Tipo"
wr["B3"] = "Tienda"
wr["C3"] = "Coordinadoras"
wr["D3"] = "Total"
for c in ("A3", "B3", "C3", "D3"):
    wr[c].font = Font(bold=True, color="FFFFFF")
    wr[c].fill = CABECERA
    wr[c].border = BORDE

for i, (t, _) in enumerate(TIPOS, 4):
    wr.cell(row=i, column=1, value=t).font = Font(bold=True)
    wr.cell(row=i, column=2, value=f'=COUNTIFS(Clasificar!$A$5:$A${ULTIMA},"Tienda",Clasificar!$E$5:$E${ULTIMA},$A{i})')
    wr.cell(row=i, column=3, value=f'=COUNTIFS(Clasificar!$A$5:$A${ULTIMA},"Coordinadoras",Clasificar!$E$5:$E${ULTIMA},$A{i})')
    wr.cell(row=i, column=4, value=f"=B{i}+C{i}")
    for col in range(1, 5):
        wr.cell(row=i, column=col).border = BORDE

fila_sin = 4 + len(TIPOS)
wr.cell(row=fila_sin, column=1, value="Sin clasificar").font = Font(bold=True, color="B3261E")
wr.cell(row=fila_sin, column=2, value=f'=COUNTIFS(Clasificar!$A$5:$A${ULTIMA},"Tienda",Clasificar!$E$5:$E${ULTIMA},"")')
wr.cell(row=fila_sin, column=3, value=f'=COUNTIFS(Clasificar!$A$5:$A${ULTIMA},"Coordinadoras",Clasificar!$E$5:$E${ULTIMA},"")')
wr.cell(row=fila_sin, column=4, value=f"=B{fila_sin}+C{fila_sin}")
for col in range(1, 5):
    wr.cell(row=fila_sin, column=col).border = BORDE

fila_tot = fila_sin + 1
wr.cell(row=fila_tot, column=1, value="TOTAL").font = Font(bold=True)
wr.cell(row=fila_tot, column=2, value=f'=COUNTIF(Clasificar!$A$5:$A${ULTIMA},"Tienda")')
wr.cell(row=fila_tot, column=3, value=f'=COUNTIF(Clasificar!$A$5:$A${ULTIMA},"Coordinadoras")')
wr.cell(row=fila_tot, column=4, value=f"=B{fila_tot}+C{fila_tot}")
for col in range(1, 5):
    wr.cell(row=fila_tot, column=col).border = BORDE
    wr.cell(row=fila_tot, column=col).fill = PatternFill("solid", fgColor=SUAVE)

wr.cell(row=fila_tot + 2, column=1,
        value="Cuando esté, mándamelo y lo meto en la aplicación. "
              "Las que se queden en «Sin clasificar» las dejo como están, sin decidir por ti.")
wr.cell(row=fila_tot + 2, column=1).font = Font(size=10, italic=True, color="5E6E63")

wr.column_dimensions["A"].width = 22
wr.column_dimensions["B"].width = 15
wr.column_dimensions["C"].width = 16
wr.column_dimensions["D"].width = 10

# ---------------------------------------------------------------- hoja: instrucciones
wi = wb.create_sheet("Cómo se hace", 0)
wi["A1"] = "CÓMO RELLENARLO"
wi["A1"].font = Font(bold=True, size=16, color="FFFFFF")
wi["A1"].fill = TITULO
wi.merge_cells("A1:B1")
wi.row_dimensions[1].height = 30

PASOS = [
    ("1", "Abre la hoja «Clasificar». Hay 93 tareas, una por fila."),
    ("2", "La columna TIPO tiene un desplegable: Rutina · Exclusiva del día · Urgente."),
    ("3", "La columna ZONA tiene otro desplegable, por si quieres decir dónde se hace."),
    ("4", "NOTAS es libre: para lo que no quepa en ninguna."),
    ("5", "En «Resumen» se ve cuántas llevas y cuántas faltan."),
    ("6", "Las filas están ordenadas poniendo primero las que tocan más días al día."),
    ("", ""),
    ("", "No hay respuestas mías en ninguna parte: los desplegables están vacíos."),
    ("", "Si una tarea no encaja en ninguno de los tres tipos, escríbelo en NOTAS y ya está."),
]
for i, (n, txt) in enumerate(PASOS, 3):
    wi.cell(row=i, column=1, value=n).font = Font(bold=True, color=VERDE_OSC)
    wi.cell(row=i, column=2, value=txt).alignment = Alignment(wrap_text=True, vertical="center")
    if txt:
        wi.row_dimensions[i].height = 22

wi.column_dimensions["A"].width = 5
wi.column_dimensions["B"].width = 96

wb.move_sheet("Cómo se hace", offset=-wb.sheetnames.index("Cómo se hace"))
wb.save("docs/clasificar-tareas.xlsx")
print(f"  escrito docs/clasificar-tareas.xlsx  ·  filas de tarea: {ULTIMA - 4}")
