"""Saca el ciclo de 4 semanas de las coordinadoras del Excel."""
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

RUTA = r"docs\referencia\horarios-verano-2026-anonimizado.xlsx"
wb = load_workbook(RUTA, data_only=True)
ws = wb["VacasTodas2026"]

DIAS = ["L", "M", "X", "J", "V", "S", "D"]
COLS = [4, 6, 8, 10, 12, 14, 16]   # D F H J L N P

ciclo = {}     # (persona, semana) -> {dia: horas}
for r in range(1, ws.max_row + 1):
    b = ws.cell(row=r, column=2).value
    q = ws.cell(row=r, column=17).value
    if not b or not isinstance(b, str):
        continue
    if "COORD" not in str(b).upper() and "Gerencia" not in str(b):
        continue
    if not q:
        continue
    marca = str(q).strip()
    horas = {}
    for i, c in enumerate(COLS):
        v = ws.cell(row=r, column=c).value
        try:
            horas[DIAS[i]] = float(v) if v is not None else 0.0
        except (TypeError, ValueError):
            horas[DIAS[i]] = 0.0
    persona = "COORD 1" if "1" in str(b) else ("COORD 2" if "2" in str(b) else str(b))
    # la primera aparición de cada marca es la buena
    if (persona, marca) not in ciclo:
        ciclo[(persona, marca)] = horas

print("  === EL CICLO, tal como está en el Excel ===")
print("  (horas de VERANO; los días con 0 no se trabaja)\n")
for persona in ("COORD 1", "COORD 2"):
    print(f"  {persona}")
    orden = ["1º", "1ª", "2º", "2ª", "3º", "3ª", "4º", "4ª"]
    for sem in ["1", "2", "3", "4"]:
        clave = next((k for k in ciclo if k[0] == persona and
                      str(k[1]).startswith(sem)), None)
        if not clave:
            continue
        h = ciclo[clave]
        total = sum(h.values())
        dias = " ".join(f"{d}:{h[d]:g}" for d in DIAS if h[d] > 0)
        libres = " ".join(d for d in DIAS if h[d] == 0)
        print(f"    semana {sem}  total {total:>5.1f} h   trabaja  {dias:<42} libre {libres}")
    print()

print("  === medias ===")
for persona in ("COORD 1", "COORD 2"):
    totales = []
    for sem in ["1", "2", "3", "4"]:
        clave = next((k for k in ciclo if k[0] == persona and
                      str(k[1]).startswith(sem)), None)
        if clave:
            totales.append(sum(ciclo[clave].values()))
    if totales:
        print(f"    {persona}: {[f'{t:g}' for t in totales]}  media {sum(totales)/len(totales):.2f} h/sem")
