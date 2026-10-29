"""Calcula el computo anual definitivo del Ejemplo de horarios 2026.

Estructura de la hoja VacasTodas2026:
  B = persona (OP 1..3, COORD 1..2, YO)
  horas por dia: D=lunes F=martes H=miercoles J=jueves L=viernes N=sabado P=domingo
  Q = marca (M/T/FD para personas de tienda, 1..4a para coordinadoras)
"""
import re
import sys
from collections import OrderedDict
from openpyxl import load_workbook

RUTA = r"C:\Users\Usuario\Desktop\APP Malato Documentación\EjemploHorariosVacacionesVerano2026.xlsx"
wb = load_workbook(RUTA, data_only=False)
wbv = load_workbook(RUTA, data_only=True)

COL_HORAS = [4, 6, 8, 10, 12, 14, 16]          # D F H J L N P
DIAS = ["Lun", "Mar", "Mie", "Jue", "Vie", "Sab", "Dom"]
PERSONAS = ["OP 1", "OP 2", "OP 3", "COORD 1", "COORD 2", "YO"]
CONTRATO_ANUAL = {"COORD 1": 1396.8, "COORD 2": 938.5,
                  "OP 1": 938.5, "OP 2": 938.5, "OP 3": 938.5}


def num(ws, wsv, r, c):
    v = ws.cell(row=r, column=c).value
    if isinstance(v, (int, float)):
        return float(v)
    v2 = wsv.cell(row=r, column=c).value
    if isinstance(v2, (int, float)):
        return float(v2)
    return None


for hoja in ["VacasTodas2026"]:
    ws, wsv = wb[hoja], wbv[hoja]
    persona = None
    semana = None
    datos = OrderedDict()
    vac_dias = {}
    libre_dias = {}
    marcas = {}
    semanas_vistas = []
    filas_con_datos = 0

    for r in range(1, ws.max_row + 1):
        et = ws.cell(row=r, column=2).value
        if et is not None:
            e = str(et).strip()
            if re.match(r"^\d{2}/\d{2}\s*-\s*\d{2}/\d{2}$", e):
                semana = e
                if e not in semanas_vistas:
                    semanas_vistas.append(e)
                continue
            if e.upper() in PERSONAS:
                persona = e.upper()
                datos.setdefault(persona, OrderedDict()).setdefault(semana, 0.0)
                marcas.setdefault(persona, {}).setdefault(semana, ws.cell(row=r, column=17).value)
                # OJO: aqui NO se hace continue. La etiqueta de persona y las horas
                # del lunes comparten fila (B4='OP 1' y D4=4). Saltarse la fila
                # descartaba las horas de todos los lunes.

        if persona is None or semana is None:
            continue

        h = sum(x for x in (num(ws, wsv, r, c) for c in COL_HORAS) if x)
        if h:
            datos[persona][semana] = round(datos[persona].get(semana, 0.0) + h, 3)
            filas_con_datos += 1
        # textos de la fila: VAC / libre y turnos
        for c in range(3, 17):
            v = ws.cell(row=r, column=c).value
            if not isinstance(v, str):
                continue
            t = v.strip().upper()
            if t == "VAC":
                vac_dias[persona] = vac_dias.get(persona, 0) + 1
            elif t == "LIBRE":
                libre_dias[persona] = libre_dias.get(persona, 0) + 1

    print("=" * 84)
    print(f"HOJA {hoja!r}")
    print("=" * 84)
    print(f"Semanas detectadas: {len(semanas_vistas)}   de {semanas_vistas[0]} a {semanas_vistas[-1]}")
    print(f"Filas de horas leidas: {filas_con_datos}\n")

    print(f"{'Persona':<9}{'Semanas':>8}{'Horas tot.':>12}{'Media/sem':>11}"
          f"{'Dias VAC':>10}{'Dias libre':>11}{'Contrato':>10}{'Desvio':>10}")
    for p, semanas in datos.items():
        vals = [v for v in semanas.values()]
        total = sum(vals)
        n = len(vals)
        media = total / n if n else 0
        ctr = CONTRATO_ANUAL.get(p)
        cel = f"{ctr:.1f}" if ctr else "autonoma"
        des = f"{total - ctr:+.1f}" if ctr else "-"
        print(f"{p:<9}{n:>8}{total:>12.1f}{media:>11.2f}"
              f"{vac_dias.get(p, 0):>10}{libre_dias.get(p, 0):>11}{cel:>10}{des:>10}")
    print()

    # detalle semanal de las primeras 6 y ultimas 3 semanas
    print("Detalle semanal (horas):")
    for p, semanas in datos.items():
        items = list(semanas.items())
        cabeza = "  ".join(f"{s.split('-')[0]}={v:g}" for s, v in items[:6])
        cola = "  ".join(f"{s.split('-')[0]}={v:g}" for s, v in items[-3:])
        print(f"  {p:<8} {cabeza}   ...   {cola}")
