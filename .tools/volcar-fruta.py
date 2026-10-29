"""Volcado completo de la clasificación, para revisar CI raros."""
import sys
from collections import defaultdict

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook("docs/fruta-verdura-familias.xlsx", data_only=True)
ws = wb["Clasificar"]

por_fam = defaultdict(list)
for r in range(5, ws.max_row + 1):
    ci = ws.cell(row=r, column=1).value
    prod = ws.cell(row=r, column=2).value
    fam = ws.cell(row=r, column=3).value
    if not prod:
        continue
    por_fam[str(fam).strip() if fam else "(sin familia)"].append(
        ("" if ci is None else str(ci).strip(), str(prod).strip()))

orden = ["Fruta", "Verdura", "Embolsado", "Ensaladas Preparadas", "Legumbres", "Varios"]
for fam in orden + [k for k in por_fam if k not in orden]:
    if fam not in por_fam:
        continue
    print(f"\n=== {fam} ({len(por_fam[fam])}) ===")
    for ci, prod in por_fam[fam]:
        marca = "  <<< CI RARO" if ci in ("96", "960156") else ""
        print(f"  {ci:<10} {prod}{marca}")
