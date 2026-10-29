"""Vuelca el contenido completo de PlantillasHorarios.xlsx en texto legible."""
import sys
from openpyxl import load_workbook

RUTA = r"C:\Users\Usuario\Desktop\APP Malato Documentación\PlantillasHorarios.xlsx"

wb = load_workbook(RUTA, data_only=False)
print(f"Hojas ({len(wb.sheetnames)}): {wb.sheetnames}\n")

for nombre in wb.sheetnames:
    ws = wb[nombre]
    print("=" * 78)
    print(f"HOJA: {nombre!r}   dimensiones={ws.dimensions}   "
          f"max_fila={ws.max_row} max_col={ws.max_column}")
    combinadas = [str(r) for r in ws.merged_cells.ranges]
    print(f"  celdas combinadas ({len(combinadas)}): {', '.join(combinadas[:20])}"
          + (" ..." if len(combinadas) > 20 else ""))
    print("=" * 78)

    vacias = 0
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        celdas = []
        for c in fila:
            v = c.value
            if v is None:
                celdas.append("")
            else:
                celdas.append(str(v).replace("\n", "\\n"))
        if not any(celdas):
            vacias += 1
            continue
        # recorta columnas vacias del final
        while celdas and celdas[-1] == "":
            celdas.pop()
        print(f"  r{fila[0].row:>3} | " + " | ".join(celdas))
    if vacias:
        print(f"  ({vacias} filas vacias omitidas)")
    print()
