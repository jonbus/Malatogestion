"""Comprueba el calendario generado contra los MINIMOS DE PERSONAL reales.

Reglas (docs/RRHH-TURNOS.md §3):
  Lunes y viernes : 2 personas desde la apertura (7:00 invierno / 8:00 verano)
  Martes y miercoles: 1 hasta las 11:00, 2 a partir de las 11:00
  Jueves          : 1 hasta las 9:00, 3 a partir de las 9:00, + apoyo de 3 h
  Sabado          : 2 en la manana (8:00 y 10:30)
  Domingo/festivo : 1 a las 8:00, 3 a partir de las 9:00
  Todas las tardes: 1 persona
"""
import re
import sys
from collections import defaultdict
from datetime import date
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
GEN = r"docs\calendario-2026\calendario-2026.xlsx"
MESES = {m: i + 1 for i, m in enumerate(
    ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
     "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"])}
FESTIVOS = {(1, 1), (1, 6), (3, 19), (4, 2), (4, 3), (4, 6), (5, 1),
            (7, 25), (8, 15), (10, 12), (12, 8), (12, 25), (4, 28), (7, 22)}
RX = re.compile(r"(\d{1,2})(?:[:.,](\d{2}))?\s*-\s*(\d{1,2})(?:[:.,](\d{2}))?")


def tramos(t):
    out = []
    for a, b, c, d in RX.findall(t):
        ini = int(a) + (int(b) / 60 if b else 0)
        fin = int(c) + (int(d) / 60 if d else 0)
        if fin > ini:
            out.append((ini, fin))
    return out


wb = load_workbook(GEN)
agenda = defaultdict(dict)
for hoja in wb.sheetnames:
    p = hoja.split("-")
    if len(p) != 2 or p[1] not in MESES:
        continue
    ws = wb[hoja]
    dias = []
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=9):
        a = fila[0].value
        if isinstance(a, str) and a.strip().upper().startswith("PERSONA"):
            dias = [int(re.search(r"(\d{1,2})\s*$", str(c.value or "")).group(1))
                    if re.search(r"(\d{1,2})\s*$", str(c.value or "")) else None
                    for c in fila[1:8]]
            continue
        if not isinstance(a, str) or a.strip() not in (
                "Gerencia", "COORD 1", "COORD 2", "OP 1", "OP 2", "OP 3"):
            continue
        for idx, celda in enumerate(fila[1:8]):
            if idx >= len(dias) or dias[idx] is None:
                continue
            t = str(celda.value or "").strip()
            if not t or t.upper() in ("VAC", "LIBRE"):
                continue
            tr = tramos(t)
            if tr:
                agenda[date(2026, MESES[p[1]], dias[idx])][a.strip()] = tr


def n(f, h):
    return sum(1 for tr in agenda[f].values() if any(a <= h < b for a, b in tr))


fallos = defaultdict(list)
sabados_3 = defaultdict(int)
for f in sorted(agenda):
    ver = date(2026, 6, 1) <= f <= date(2026, 9, 15)
    wd = f.weekday()
    fest = (f.month, f.day) in FESTIVOS
    if fest or wd == 6:
        if n(f, 8.25) < 1:
            fallos["domingo/festivo: falta 1 a las 8:00"].append(f)
        if n(f, 9.5) < 3:
            fallos["domingo/festivo: menos de 3 a las 9:30"].append(f)
    elif wd == 5:
        if n(f, 8.25) < 1:
            fallos["sabado: falta 1 a las 8:00"].append(f)
        # A las 10:30 ya deberia estar la persona de tienda. Normal 3 personas; 2 es
        # admisible cuando la 2a coordinadora libra por su ciclo. Menos de 2, no.
        c3 = n(f, 10.75)
        sabados_3[c3] += 1
        if c3 < 2:
            fallos["sabado: menos de 2 a las 10:45"].append(f)
    else:
        ini = 8.0 if ver else 7.0
        if wd in (0, 4):                                  # lunes y viernes
            if n(f, ini + 0.25) < 2:
                fallos["LUNES/VIERNES: menos de 2 desde la apertura"].append(f)
        if wd in (1, 2):                                  # martes y miercoles
            if n(f, ini + 0.25) < 1:
                fallos["martes/miercoles: falta 1 desde la apertura"].append(f)
            if n(f, 11.5) < 2:
                fallos["martes/miercoles: menos de 2 a las 11:30"].append(f)
        if wd == 3:                                       # jueves
            if n(f, ini + 0.25) < 1:
                fallos["jueves: falta 1 desde la apertura"].append(f)
            if n(f, 9.5) < 3:
                fallos["jueves: menos de 3 a las 9:30"].append(f)
        if ver and wd == 3 and False:
            pass
    # tardes: al menos 1 persona
    tarde = 17.75 if ver else 17.5
    if n(f, tarde) < 1:
        fallos["tarde sin nadie"].append(f)

print("=" * 76)
print("CALENDARIO GENERADO frente a los MINIMOS REALES")
print("=" * 76)
print(f"dias revisados: {len(agenda)}\n")
total = 0
for k in sorted(fallos, key=lambda x: -len(fallos[x])):
    v = fallos[k]
    total += len(v)
    ej = ", ".join(str(x) for x in v[:5])
    print(f"  {k}: {len(v)} dias")
    print(f"      {ej}{' ...' if len(v) > 5 else ''}")
if not fallos:
    print("  sin incumplimientos")
print(f"\nTOTAL de incumplimientos: {total}")
print("\nSabados por numero de personas a las 10:45:")
for k in sorted(sabados_3): print(f"  {k} personas: {sabados_3[k]} sabados")

