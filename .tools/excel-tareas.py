"""Excel con TODAS las tareas y de quién son, para revisar las cascadas.

Sale de los datos de la propia aplicación, así que no se queda desfasado.
"""
import json
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

h = open("docs/nube/public/index.html", encoding="utf-8").read()
js = h.split("<script>")[-1].split("</script>")[0]

i = js.find("const DATOS = ")
prof = 0
for k in range(js.find("{", i), len(js)):
    if js[k] == "{":
        prof += 1
    elif js[k] == "}":
        prof -= 1
        if prof == 0:
            DATOS = json.loads(js[js.find("{", i):k + 1])
            break

NOMBRE = {"tienda": "Tienda", "coordinacion": "Coordinación", "gerencia": "Jefa"}
MOMENTO = {
    "manana": "Mañana", "tarde": "Tarde", "cierre": "Cierre",
    "apertura": "Apertura", "dia": "El día",
}

filas = []


def mete(puesto, momento, bloque, tarea):
    partes = tarea.get("partes") or []
    filas.append({
        "puesto": NOMBRE.get(puesto, puesto),
        "momento": momento,
        "bloque": bloque,
        "tarea": tarea.get("t"),
        "nota": tarea.get("n") or "",
        "partes": " · ".join(partes) if partes else "",
        "cuantas": len(partes) if partes else 1,
    })


# --- tienda
for momento, dato in (DATOS.get("tienda") or {}).items():
    if momento == "productos":
        continue
    if isinstance(dato, dict):
        for sec in ("rutina", "antes", "porDia"):
            if sec == "porDia":
                for dia, lista in (dato.get("porDia") or {}).items():
                    for x in lista:
                        mete("tienda", momento, "Lo de hoy · " + dia, x)
            else:
                for x in (dato.get(sec) or []):
                    nombre_bloque = {"rutina": "Lo de siempre", "antes": "Antes de irse"}[sec]
                    mete("tienda", momento, nombre_bloque, x)
    elif isinstance(dato, list):
        for sec in dato:
            for x in (sec.get("tareas") or []):
                mete("tienda", momento, sec.get("sec", ""), x)

# --- coordinación
co = DATOS.get("coordinacion") or {}
for x in (co.get("dia") or {}).get("rutina") or []:
    mete("coordinacion", "El día", "Rutina · todos los días", x)
for dia, lista in ((co.get("dia") or {}).get("porDia") or {}).items():
    for x in lista:
        mete("coordinacion", "El día", "Lo de hoy · " + dia, x)
for pat, dato in (co.get("apertura") or {}).items():
    if isinstance(dato, dict):
        for b in (dato.get("bloques") or []):
            for x in (b.get("tareas") or []):
                mete("coordinacion", "Apertura", b.get("sec", ""), x)
for sec in (co.get("cierre") or []):
    if isinstance(sec, dict):
        for x in (sec.get("tareas") or []):
            mete("coordinacion", "Cierre", sec.get("sec", ""), x)

# --- gerencia
ge = DATOS.get("gerencia") or {}
for clave in ("lunes", "martes", "miercoles", "jueves", "viernes", "sabado", "domingo"):
    for x in (ge.get(clave) or []):
        if isinstance(x, dict):
            mete("gerencia", "El día", clave.capitalize(), x)
for x in (ge.get("hoy") or []):
    mete("gerencia", "El día", "Lo de cada día", x)

print(f"  tareas recogidas: {len(filas)}")

wb = Workbook()
ws = wb.active
ws.title = "Todas las tareas"

cab = ["Puesto", "Momento", "Bloque", "Tarea", "Nota de la propiedad",
       "¿Partes?", "Nº partes", "¿Es cascada?", "Después va a", "Puesto del siguiente"]
ws.append(cab)
for f in filas:
    ws.append([f["puesto"], f["momento"], f["bloque"], f["tarea"], f["nota"],
               "Sí" if f["cuantas"] > 1 else "No", f["cuantas"], "", "", ""])

# estilo
verde = PatternFill("solid", fgColor="2E7D32")
for c in range(1, len(cab) + 1):
    cel = ws.cell(row=1, column=c)
    cel.font = Font(bold=True, color="FFFFFF")
    cel.fill = verde
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

anchos = [14, 12, 26, 62, 34, 10, 9, 12, 30, 18]
for i, a in enumerate(anchos, start=1):
    ws.column_dimensions[get_column_letter(i)].width = a

for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cel in fila:
        cel.alignment = Alignment(vertical="top", wrap_text=True)
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{ws.max_row}"

# --- hoja de resumen
ws2 = wb.create_sheet("Cómo va")
ws2.append(["Puesto", "Tareas"])
ws2["A1"].font = Font(bold=True)
ws2["B1"].font = Font(bold=True)
porPuesto = {}
for f in filas:
    porPuesto[f["puesto"]] = porPuesto.get(f["puesto"], 0) + 1
for k, v in sorted(porPuesto.items()):
    ws2.append([k, v])
ws2.append(["TOTAL", len(filas)])
ws2.column_dimensions["A"].width = 20
ws2.column_dimensions["B"].width = 12

salida = "TAREAS-APP-MALATO.xlsx"
wb.save(salida)
print(f"  escrito {salida}")
print(f"  puestos: {json.dumps(porPuesto, ensure_ascii=False)}")
