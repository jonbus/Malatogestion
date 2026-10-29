"""Analiza las plantillas de horarios.

Corrige el error de la primera version: la hoja de encargadas repite los mismos
bloques de semanas, y el script sumaba las repeticiones como si fueran semanas
distintas. Ahora cada (persona, semana) se cuenta UNA sola vez: la primera
aparicion. Ademas informa de los colores de celda, que suelen marcar festivos.
"""
import re
from collections import OrderedDict
from openpyxl import load_workbook

RUTA = r"C:\Users\Usuario\Desktop\APP Malato Documentación\PlantillasHorarios.xlsx"
CONTRATO = {"COORD 1": 32.0, "COORD2": 21.5, "OP 1": 21.5, "OP 2": 21.5, "OP 3": 21.5,
            "GER": None, "YO": None}
DIAS = {"LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"}
wb = load_workbook(RUTA, data_only=True)


def a_horas(t):
    t = str(t).strip().replace(" ", "")
    m = re.match(r"^(\d{1,2})(?:[,.](\d{2}))?-(\d{1,2})(?:[,.](\d{2}))?$", t)
    if not m:
        return None
    h1, m1, h2, m2 = m.groups()
    ini = int(h1) + (int(m1) / 60 if m1 else 0)
    fin = int(h2) + (int(m2) / 60 if m2 else 0)
    return round(fin - ini, 3)


def semana_idx(e):
    """'3ª Semana' y '3' son la misma semana: se normalizan al digito."""
    m = re.search(r"\d", str(e))
    return m.group(0) if m else str(e)


def color(celda):
    f = celda.fill
    if f is None or f.fill_type is None:
        return None
    rgb = getattr(f.start_color, "rgb", None)
    if not rgb or rgb in ("00000000",):
        return None
    return str(rgb)[-6:]


for hoja in ["Turnos Gerencia", "Turnos EncargadasIndividualSinC", "Turnos Tienda"]:
    ws = wb[hoja]
    print("=" * 78)
    print(f"HOJA {hoja!r}")
    print("=" * 78)

    persona = None
    semana = None
    datos = OrderedDict()      # persona -> OrderedDict(semana -> horas)
    vistos = set()             # (persona, semana) ya procesados
    duplicado = False
    colores = {}
    raros = []

    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=8):
        et = fila[0].value

        # Fila de cabecera de semana (B..H son nombres de dia): no es un turno.
        if str(fila[1].value or "").strip().upper() in DIAS:
            if et is not None and ("semana" in str(et).lower() or re.match(r"^\d$", str(et).strip())):
                semana = semana_idx(et)
            continue

        if et is not None:
            e = str(et).strip()
            if "semana" in e.lower() or re.match(r"^\d$", e):
                semana = semana_idx(e)
            elif e.upper() in CONTRATO:
                persona = e.upper()
                clave = (persona, semana)
                # La hoja repite bloques con la MISMA etiqueta de semana: a partir
                # de la segunda vez son copias, no semanas nuevas.
                duplicado = clave in vistos
                vistos.add(clave)
                if not duplicado:
                    datos.setdefault(persona, OrderedDict()).setdefault(semana, 0.0)

        if persona is None or semana is None or duplicado:
            continue

        for c in fila[1:8]:
            if c.value is None or str(c.value).strip() == "":
                continue
            h = a_horas(c.value)
            if h is None:
                raros.append((c.row, persona, str(c.value)))
                continue
            datos[persona][semana] = round(datos[persona][semana] + h, 3)
            col = color(c)
            if col:
                colores.setdefault(col, []).append(f"{persona} {semana} {c.value}")

    for p, semanas in datos.items():
        vals = list(semanas.values())
        media = sum(vals) / len(vals) if vals else 0
        ctr = CONTRATO.get(p)
        det = "  ".join(f"{s[:3]}={v:g}" for s, v in semanas.items())
        print(f"  {p:<8} {det}")
        if ctr:
            print(f"           media {media:.2f} h/semana  |  contrato {ctr:g}h  |  "
                  f"DESVIO {media - ctr:+.2f} h/semana  "
                  f"({(media - ctr) * 52:+.0f} h/año)")
        else:
            print(f"           media {media:.2f} h/semana  |  autonomia, sin tope")
    if raros:
        print(f"  celdas no interpretables: {raros[:6]}")
    if colores:
        print("  COLORES DE CELDA (posible codificacion de festivos):")
        for col, usos in colores.items():
            print(f"    #{col}: {len(usos)} celdas -> {usos[:5]}")
    else:
        print("  sin colores de relleno")
    print()

