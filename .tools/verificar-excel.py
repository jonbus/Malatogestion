"""Verifica el contenido real del Excel de revisión."""
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook("docs/tareas-coordinadora.xlsx")
print("  hojas:", wb.sheetnames)
print()
for n in wb.sheetnames:
    ws = wb[n]
    if n == "Cómo usarlo":
        continue
    tareas = 0
    secciones = 0
    for r in range(5, ws.max_row + 1):
        c1 = ws.cell(row=r, column=1).value
        c2 = ws.cell(row=r, column=2).value
        if not c1 and not c2:
            continue
        # seccion = fila fusionada o sin la otra columna
        if (c1 and not c2) or (c2 and not c1):
            f = ws.cell(row=r, column=1).font
            if f and f.bold and f.color and f.color.rgb and "1E9A45" in str(f.color.rgb):
                secciones += 1
                continue
        tareas += 1
    print(f"  {n:<24} tareas={tareas:<3} apartados={secciones}")

print("\n  === PRIMERAS FILAS DE «Tienda mañana» ===")
ws = wb["Tienda mañana"]
for r in range(5, 16):
    c1 = ws.cell(row=r, column=1).value
    c2 = ws.cell(row=r, column=2).value
    if c1 or c2:
        print(f"    r{r}: {(c1 or '')[:40]:<42} | {(c2 or '')[:52]}")

print("\n  === PRIMERAS FILAS DE «Tienda tarde» ===")
ws = wb["Tienda tarde"]
for r in range(5, 14):
    c1 = ws.cell(row=r, column=1).value
    c2 = ws.cell(row=r, column=2).value
    if c1 or c2:
        print(f"    r{r}: {(c1 or '')[:40]:<42} | {(c2 or '')[:52]}")
