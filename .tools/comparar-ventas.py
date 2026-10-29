"""Compara 2026 con 2025 bloque a bloque (mismo numero de semana)."""
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook(r".tools\ventas-pollos.xlsx", data_only=True)


def numero(v, defecto=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return defecto


def bloques(hoja, producto="POLLOS"):
    """Devuelve la lista de totales semanales del producto, en orden."""
    ws = wb[hoja]
    salida = []
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip().upper() == producto:
            salida.append(numero(ws.cell(row=r, column=14).value))
    return salida


a25 = bloques("2025")
a26 = bloques("2026")
print(f"Bloques con fila de POLLOS: 2025={len(a25)}  2026={len(a26)}")

# Solo los bloques de 2026 que tienen algo vendido (el año esta a medias)
con_datos = [i for i, v in enumerate(a26) if v > 0]
print(f"Bloques de 2026 con venta registrada: {len(con_datos)}")

t25 = sum(a25[i] for i in con_datos[:0]) if False else 0
# Comparacion semana a semana sobre los bloques que 2026 tiene con datos
print("\n  semana   2025   2026   variacion")
t25 = t26 = 0.0
for i in con_datos:
    v25, v26 = a25[i], a26[i]
    t25 += v25
    t26 += v26
    var = (v26 / v25 - 1) * 100 if v25 else float("nan")
    print(f"  {i + 1:>5}  {v25:>6.0f} {v26:>6.0f}   {var:>+7.1f} %")
print(f"  {'TOTAL':>5}  {t25:>6.0f} {t26:>6.0f}   {100 * (t26 / t25 - 1):>+7.1f} %")

print("\n  Comparacion con el AÑO COMPLETO anterior:")
print(f"    2025 completo: {sum(a25):.0f} pollos")
print(f"    2024 completo: {sum(bloques('2024')):.0f} pollos")
print(f"    2023 completo: {sum(bloques('2023')):.0f} pollos")
