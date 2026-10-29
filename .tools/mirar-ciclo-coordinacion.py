"""Busca el ciclo de las coordinadoras en el Excel de referencia."""
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

for ruta in ("docs/referencia/horarios-verano-2026-anonimizado.xlsx",):
    wb = load_workbook(ruta, data_only=True)
    print(f"\n=== {ruta} ===")
    print(f"  hojas: {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    print(f"  filas: {ws.max_row}  columnas: {ws.max_column}")

    # buscar las filas de COORD
    for r in range(1, min(ws.max_row, 120) + 1):
        b = ws.cell(row=r, column=2).value
        q = ws.cell(row=r, column=17).value
        if b and isinstance(b, str) and ("COORD" in str(b).upper() or "Gerencia" in str(b)):
            horas = [ws.cell(row=r, column=c).value for c in (4, 6, 8, 10, 12, 14, 16)]
            print(f"    r{r:<4} {str(b):<14} marca={str(q):<5} "
                  f"L-V-S-D={[str(h) if h is not None else '-' for h in horas]}")
