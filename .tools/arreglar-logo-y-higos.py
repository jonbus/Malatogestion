"""Pone el logo en las dos pantallas del archivo y completa el nombre de los higos."""
import json
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

# ---------------------------------------------------- 1) el logo que faltaba
HUEco = '<div class="logo">AQUÍ<br>EL LOGO</div>'
NUEVO = ('<div class="logo"><img src="../marca/logo-o-blanco.png" alt="" '
         'style="max-width:100%;max-height:100%;object-fit:contain"></div>')

for ruta in ("docs/archivo/mi-dia.html", "docs/archivo/vista-tareas.html"):
    t = open(ruta, encoding="utf-8").read()
    if HUEco in t:
        t = t.replace(HUEco, NUEVO)
        # el hueco era un cuadro con borde discontinuo; ahora es una imagen
        t = t.replace("border:2px dashed var(--verde-claro)", "border:none")
        open(ruta, "w", encoding="utf-8").write(t)
        print(f"  logo puesto en {ruta}")
    else:
        print(f"  no encuentro el hueco en {ruta}")

# ---------------------------------------------------- 2) los higos
RUTA_XLSX = "docs/fruta-verdura-familias.xlsx"
wb = load_workbook(RUTA_XLSX)
ws = wb["Clasificar"]

cambiada = False
for r in range(5, ws.max_row + 1):
    if str(ws.cell(row=r, column=1).value).strip() == "960050":
        viejo = ws.cell(row=r, column=2).value
        nuevo = "HIGOS CUELLO DAMA BOLSA 500GR"
        ws.cell(row=r, column=2, value=nuevo)
        print(f"  fila {r}: «{viejo}» -> «{nuevo}»")
        cambiada = True
        break

if cambiada:
    wb.save(RUTA_XLSX)
    print("  Excel guardado")
else:
    print("  no encuentro el 960050")

# ---------------------------------------------------- 3) el JSON de la app
wb2 = load_workbook(RUTA_XLSX, data_only=True)
ws2 = wb2["Clasificar"]
FAMILIAS = ["Fruta", "Verdura", "Embolsado", "Ensaladas Preparadas", "Legumbres", "Varios"]
productos = []
for r in range(5, ws2.max_row + 1):
    ci = ws2.cell(row=r, column=1).value
    prod = ws2.cell(row=r, column=2).value
    fam = ws2.cell(row=r, column=3).value
    if not prod:
        continue
    productos.append({
        "ci": "" if ci is None else str(ci).strip(),
        "nombre": str(prod).strip(),
        "familia": str(fam).strip() if fam else "",
    })

datos = {
    "generado_de": "docs/fruta-verdura-familias.xlsx (completado por la propiedad)",
    "familias": FAMILIAS,
    "aviso_ci": "El CI es interno: solo gerencia y coordinadoras. Sirve para los pedidos.",
    "productos": productos,
}
with open("docs/datos/fruta-verdura.json", "w", encoding="utf-8") as f:
    json.dump(datos, f, ensure_ascii=False, indent=2)

from collections import Counter
c = Counter(p["familia"] for p in productos)
print(f"\n  JSON regenerado: {len(productos)} productos")
for fam in FAMILIAS:
    print(f"    {fam:<24} {c[fam]:>3}")
