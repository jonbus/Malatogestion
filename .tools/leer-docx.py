"""Lee un .docx y volca su contenido: parrafos, listas y tablas, en orden."""
import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
RUTA = sys.argv[1]
doc = Document(RUTA)

print("=" * 78)
print("CONTENIDO DEL DOCUMENTO")
print("=" * 78)

# Recorrer el cuerpo en orden, mezclando parrafos y tablas
from docx.oxml.ns import qn

cuerpo = doc.element.body
tablas_por_elemento = {t._element: t for t in doc.tables}
parrafos_por_elemento = {p._element: p for p in doc.paragraphs}

for el in cuerpo.iterchildren():
    if el in parrafos_por_elemento:
        p = parrafos_por_elemento[el]
        t = p.text.strip()
        if not t:
            continue
        estilo = (p.style.name or "").lower()
        prefijo = ""
        if "heading" in estilo or "título" in estilo or "titulo" in estilo:
            prefijo = "# "
        elif "list" in estilo:
            prefijo = "- "
        print(f"{prefijo}{t}")
    elif el in tablas_por_elemento:
        tb = tablas_por_elemento[el]
        print(f"\n  [TABLA {len(tb.rows)}x{len(tb.columns)}]")
        for fila in tb.rows:
            celdas = [c.text.replace("\n", " / ").strip() for c in fila.cells]
            while celdas and celdas[-1] == "":
                celdas.pop()
            if any(celdas):
                print("    | " + " | ".join(celdas))
