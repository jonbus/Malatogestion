"""Computo por persona y semana sobre un rango de filas, contando cada dia UNA vez.

Estructura: B=persona; por dia el par (turno, horas):
  C/D lunes - E/F martes - G/H miercoles - I/J jueves - K/L viernes - M/N sabado - O/P domingo
Cada persona ocupa 2 filas (el turno puede estar repartido entre ambas).
"""
import re
import sys
from collections import OrderedDict
from openpyxl import load_workbook

RUTA = r"C:\Users\Usuario\Desktop\APP Malato Documentación\EjemploHorariosVacacionesVerano2026.xlsx"
DESDE, HASTA = int(sys.argv[1]), int(sys.argv[2])
ETIQUETA = sys.argv[3]

PARES = [(3, 4), (5, 6), (7, 8), (9, 10), (11, 12), (13, 14), (15, 16)]
DIAS = ["Lun", "Mar", "Mie", "Jue", "Vie", "Sab", "Dom"]

wb = load_workbook(RUTA, data_only=False)
wsv = load_workbook(RUTA, data_only=True)
ws, wsv = wb["VacasTodas2026"], wsv["VacasTodas2026"]


def es_persona(v):
    if not isinstance(v, str):
        return False
    s = v.strip()
    if re.match(r"^\d{2}/\d{2}\s*-\s*\d{2}/\d{2}$", s):
        return False
    if "semana" in s.lower():
        return False
    return bool(re.match(r"^[A-Za-zÁÉÍÓÚÑáéíóúñ][A-Za-z0-9ÁÉÍÓÚÑáéíóúñ.\s]{1,14}$", s)) and len(s) >= 2


def valor(r, c):
    v = ws.cell(row=r, column=c).value
    if isinstance(v, (int, float)):
        return float(v)
    v2 = wsv.cell(row=r, column=c).value
    return float(v2) if isinstance(v2, (int, float)) else None


# 1) Localizar bloques: (persona, semana, fila_inicio, fila_fin)
bloques = []
semana = None
for r in range(DESDE, HASTA + 1):
    b = ws.cell(row=r, column=2).value
    if isinstance(b, str) and re.match(r"^\d{2}/\d{2}\s*-\s*\d{2}/\d{2}$", b.strip()):
        semana = b.strip()
        continue
    if es_persona(b):
        if bloques:
            bloques[-1][3] = r - 1
        bloques.append([b.strip().upper(), semana, r, HASTA])

print(f"Seccion {ETIQUETA}: filas {DESDE}-{HASTA}  ·  bloques de persona: {len(bloques)}\n")

resumen = OrderedDict()
detalle = OrderedDict()

for persona, sem, ini, fin in bloques:
    tot_h = 0.0
    vac = libre = 0
    for tcol, hcol in PARES:
        h = 0.0
        textos = []
        for r in range(ini, fin + 1):
            v = valor(r, hcol)
            if v:
                h += v
            for c in (tcol, hcol):
                x = ws.cell(row=r, column=c).value
                if isinstance(x, str):
                    textos.append(x.strip())
            # fin de bloque de 2 filas
            if r > ini and ws.cell(row=r + 1, column=2).value is not None:
                break
        tot_h += h
        t = " ".join(textos).upper()
        if "VAC" in t:
            vac += 1
        elif "LIBRE" in t:
            libre += 1
    resumen.setdefault(persona, {"h": 0.0, "semanas": set(), "vac": 0, "libre": 0})
    resumen[persona]["h"] += tot_h
    resumen[persona]["semanas"].add(sem)
    resumen[persona]["vac"] += vac
    resumen[persona]["libre"] += libre
    detalle.setdefault(persona, OrderedDict())[sem] = round(tot_h, 2)

print(f"{'Persona':<10}{'Semanas':>8}{'Horas':>9}{'Media/sem':>11}{'Dias VAC':>10}{'Dias libre':>11}")
for p, d in resumen.items():
    n = len(d["semanas"])
    print(f"{p:<10}{n:>8}{d['h']:>9.1f}{d['h'] / n if n else 0:>11.2f}"
          f"{d['vac']:>10}{d['libre']:>11}")

print("\nDetalle semanal (horas):")
for p, semanas in detalle.items():
    print(f"  {p:<10} " + "  ".join(f"{s.split('-')[0]}={v:g}" for s, v in semanas.items()))
