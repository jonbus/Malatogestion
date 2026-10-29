"""Vuelca cualquier libro de horarios en texto legible, con formulas visibles."""
import sys
from openpyxl import load_workbook

RUTA = sys.argv[1]
wb = load_workbook(RUTA, data_only=False)
print(f"FICHERO: {RUTA}")
print(f"Hojas ({len(wb.sheetnames)}): {wb.sheetnames}\n")

for nombre in wb.sheetnames:
    ws = wb[nombre]
    print("=" * 90)
    print(f"HOJA: {nombre!r}  dims={ws.dimensions}  max_fila={ws.max_row} max_col={ws.max_column}")
    comb = [str(r) for r in ws.merged_cells.ranges]
    print(f"  combinadas ({len(comb)}): {', '.join(comb[:25])}{' ...' if len(comb) > 25 else ''}")
    print("=" * 90)
    vacias = 0
    for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200), max_col=min(ws.max_column, 20)):
        celdas = []
        for c in fila:
            v = c.value
            celdas.append("" if v is None else str(v).replace("\n", "\\n"))
        if not any(celdas):
            vacias += 1
            continue
        while celdas and celdas[-1] == "":
            celdas.pop()
        print(f"  r{fila[0].row:>3} | " + " | ".join(celdas))
    if vacias:
        print(f"  ({vacias} filas vacias omitidas)")
    if ws.max_row > 200:
        print(f"  (truncado en la fila 200 de {ws.max_row})")
    print()
