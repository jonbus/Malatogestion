"""Genera una copia anonimizada del libro de horarios.

Sustituye UNICAMENTE los nombres propios por etiquetas genericas.
No toca numeros, ni formulas, ni formato.

Al terminar comprueba celda a celda que lo unico que ha cambiado son las etiquetas.
No imprime ningun nombre en ningun momento.
"""
import os
import shutil
from openpyxl import load_workbook

ORIGEN = r"C:\Users\Usuario\Desktop\APP Malato Documentación\EjemploHorariosVacacionesVerano2026.xlsx"
DESTINO = r"C:\Users\Usuario\Desktop\APP Malato\docs\referencia\horarios-verano-2026-anonimizado.xlsx"

HOJA = "VacasTodas2026"

# Filas donde el libro tiene los nombres reales, en orden de aparicion.
FILAS_TIENDAS = [438, 440, 442]
FILAS_COORDINADORAS = [447, 449]
ETIQUETAS = ["OP 1", "OP 2", "OP 3", "COORD 1", "COORD 2"]

wb = load_workbook(ORIGEN, data_only=False)
ws = wb[HOJA]

# 1) Construir el mapa leyendo la estructura, sin escribir los nombres en pantalla.
mapa = {}
for fila, etiqueta in zip(FILAS_TIENDAS + FILAS_COORDINADORAS, ETIQUETAS):
    v = ws.cell(row=fila, column=2).value
    if isinstance(v, str) and v.strip():
        mapa[v.strip()] = etiqueta

print(f"Nombres distintos a sustituir: {len(mapa)}")
print(f"Etiquetas asignadas: {sorted(mapa.values())}")

# 2) Inventario previo de elementos que openpyxl podria no conservar.
print("\nElementos graficos del original:")
for nombre in wb.sheetnames:
    s = wb[nombre]
    print(f"  {nombre}: imagenes={len(getattr(s, '_images', []))} "
          f"graficos={len(getattr(s, '_charts', []))}")

# 3) Sustituir en TODAS las hojas y celdas.
cambios = 0
celdas_cambiadas = set()
for nombre in wb.sheetnames:
    s = wb[nombre]
    for fila in s.iter_rows():
        for c in fila:
            v = c.value
            if not isinstance(v, str):
                continue
            nuevo = v
            for real, generico in mapa.items():
                if real in nuevo:
                    nuevo = nuevo.replace(real, generico)
            if nuevo != v:
                c.value = nuevo
                cambios += 1
                celdas_cambiadas.add((nombre, c.coordinate))

print(f"\nCeldas con etiqueta sustituida: {cambios}")

os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
wb.save(DESTINO)
print(f"Guardado en: {os.path.relpath(DESTINO, r'C:\Users\Usuario\Desktop\APP Malato')}")

# 4) VERIFICACION A: no queda ningun nombre.
comprobado = load_workbook(DESTINO, data_only=False)
restos = 0
for nombre in comprobado.sheetnames:
    for fila in comprobado[nombre].iter_rows():
        for c in fila:
            if isinstance(c.value, str):
                for real in mapa:
                    if real in c.value:
                        restos += 1
print(f"\nVERIFICACION A · nombres que quedan: {restos}")

# 5) VERIFICACION B: todo lo demas es identico, celda a celda.
orig = load_workbook(ORIGEN, data_only=False)
diferencias_fuera = 0
comparadas = 0
for nombre in orig.sheetnames:
    so, sc = orig[nombre], comprobado[nombre]
    maxf = max(so.max_row, sc.max_row)
    maxc = max(so.max_column, sc.max_column)
    for r in range(1, maxf + 1):
        for col in range(1, maxc + 1):
            a = so.cell(row=r, column=col).value
            b = sc.cell(row=r, column=col).value
            comparadas += 1
            if a == b:
                continue
            if (nombre, sc.cell(row=r, column=col).coordinate) in celdas_cambiadas:
                continue          # cambio autorizado: era una etiqueta
            diferencias_fuera += 1
            if diferencias_fuera <= 5:
                print(f"  DIFERENCIA NO AUTORIZADA en {nombre}!{sc.cell(row=r, column=col).coordinate}")

print(f"VERIFICACION B · celdas comparadas: {comparadas} · "
      f"diferencias fuera de las etiquetas: {diferencias_fuera}")
print("\nRESULTADO:", "OK" if (restos == 0 and diferencias_fuera == 0) else "REVISAR")
