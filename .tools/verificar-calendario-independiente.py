"""Verificacion INDEPENDIENTE del calendario 2026.

No usa el generador ni el JSON ni la columna HORAS del libro: recalcula todo
leyendo el TEXTO de los turnos del Excel y aplicando las reglas a mano.
Si los numeros coinciden por dos caminos distintos, la entrega es solida.
"""
import re
import sys
from collections import defaultdict
from datetime import date, timedelta
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

RUTA = r"docs\calendario-2026\calendario-2026.xlsx"
CONTRATO = {"COORD 1": 1396.8, "COORD 2": 938.5, "OP 1": 938.5, "OP 2": 938.5, "OP 3": 938.5}
MESES = {m: i + 1 for i, m in enumerate(
    ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
     "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"])}

# Festivos: 12 confirmados + los 2 locales que el generador marca como SUPUESTO
FESTIVOS = {(1, 1), (1, 6), (3, 19), (4, 2), (4, 3), (4, 6), (5, 1),
            (7, 25), (8, 15), (10, 12), (12, 8), (12, 25), (4, 28), (7, 22)}

wb = load_workbook(RUTA)


def duracion(texto):
    """Suma las horas de todos los rangos horarios del texto de una celda."""
    total = 0.0
    for h1, m1, h2, m2 in re.findall(r"(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})", texto):
        ini = int(h1) + int(m1) / 60
        fin = int(h2) + int(m2) / 60
        if fin < ini:
            fin += 24          # turno que cruza medianoche
        total += fin - ini
    return round(total, 3)


def rangos(texto):
    salida = []
    for h1, m1, h2, m2 in re.findall(r"(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})", texto):
        ini = int(h1) + int(m1) / 60
        fin = int(h2) + int(m2) / 60
        salida.append((ini, fin))
    return salida


# ---------------------------------------------------------------- lectura
turnos = defaultdict(list)      # fecha -> [(persona, texto, [(ini,fin)])]
horas = defaultdict(float)      # persona -> horas
cocinas = defaultdict(int)      # persona -> dias de cocina
dias_trabajados = defaultdict(int)

for hoja in wb.sheetnames:
    partes = hoja.split("-")
    if len(partes) != 2 or partes[1] not in MESES:
        continue
    mes = MESES[partes[1]]
    ws = wb[hoja]
    dia_actuales = []
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=9):
        a = fila[0].value
        if isinstance(a, str) and a.strip().upper().startswith("PERSONA"):
            dia_actuales = []
            for c in fila[1:8]:
                m = re.search(r"(\d{1,2})\s*$", str(c.value or ""))
                dia_actuales.append(int(m.group(1)) if m else None)
            continue
        if not isinstance(a, str) or a.strip() not in CONTRATO and a.strip() != "Gerencia":
            continue
        persona = a.strip()
        for idx, celda in enumerate(fila[1:8]):
            if idx >= len(dia_actuales) or dia_actuales[idx] is None:
                continue
            texto = str(celda.value or "").strip()
            if not texto or texto.upper() in ("VAC", "LIBRE", "-"):
                continue
            d = date(2026, mes, dia_actuales[idx])
            turnos[d].append((persona, texto, rangos(texto)))
            horas[persona] += duracion(texto)
            dias_trabajados[persona] += 1
            if "COCINA" in texto.upper():
                cocinas[persona] += 1

print("=" * 74)
print("VERIFICACION INDEPENDIENTE · leyendo el texto de los turnos del Excel")
print("=" * 74)
print(f"Dias con turnos asignados: {len(turnos)}")

print("\n[1] HORAS ANUALES (recalculadas desde el texto de cada turno)")
for p in ["COORD 1", "COORD 2", "OP 1", "OP 2", "OP 3", "Gerencia"]:
    h = horas[p]
    c = CONTRATO.get(p)
    if c:
        des = h - c
        pct = des / c * 100
        tope = c * 0.05
        print(f"  {p:<9} {h:>8.2f} h | contrato {c:>7.1f} | desvio {des:>+7.2f} "
              f"({pct:>+6.2f} %) | margen 5 % {tope:>5.1f} | "
              f"{'OK' if abs(des) <= tope else 'FUERA DE MARGEN'}")
    else:
        print(f"  {p:<9} {h:>8.2f} h | autonomia, sin tope")

print("\n[2] COBERTURA por franjas (recuento de personas presentes)")
incumplimientos = []
for d in sorted(turnos):
    inv = date(2026, 6, 1) <= d <= date(2026, 9, 15)
    finde = d.weekday() == 6 or (d.month, d.day) in FESTIVOS
    sabado = d.weekday() == 5 and (d.month, d.day) not in FESTIVOS
    jueves = d.weekday() == 3 and not finde

    if finde:
        m_ini, m_fin, t_ini = 10.0, 14.5, 17.0
        minimo_manana = 3
    elif sabado:
        m_ini, m_fin = (9.0, 14.5) if inv else (9.0, 14.5)
        t_ini = 17.5 if inv else 17.0
        minimo_manana = 2
    else:
        m_ini, m_fin = (9.0, 14.5) if inv else (8.0, 14.5)
        t_ini = 17.5 if inv else 17.0
        minimo_manana = 3 if jueves else 2

    punto = min(m_ini + 2, m_fin - 0.5)
    pm = sum(1 for _, _, rs in turnos[d] if any(a <= punto < b for a, b in rs))
    pt = sum(1 for _, _, rs in turnos[d] if any(a <= t_ini + 1 < b for a, b in rs))
    if pm < minimo_manana:
        incumplimientos.append((d, "manana", pm, minimo_manana))
    if pt < 1:
        incumplimientos.append((d, "tarde", pt, 1))

print(f"  dias revisados: {len(turnos)}")
print(f"  incumplimientos de minimo: {len(incumplimientos)}")
for d, franja, hay, pide in incumplimientos[:10]:
    print(f"    {d} {franja}: hay {hay}, pide {pide}")

print("\n[3] DOMINGOS Y FESTIVOS: personas de tienda presentes")
opera = ["OP 1", "OP 2", "OP 3"]
mal_dom = 0
revisados = 0
for d in sorted(turnos):
    if d.weekday() != 6 and (d.month, d.day) not in FESTIVOS:
        continue
    revisados += 1
    n = sum(1 for p, _, _ in turnos[d] if p in opera)
    if n != 1:
        mal_dom += 1
        if mal_dom <= 6:
            print(f"    {d}: {n} personas de tienda (se espera exactamente 1)")
print(f"  domingos+festivos revisados: {revisados} | con exactamente 1 persona de tienda: {revisados - mal_dom}")

print("\n[4] COCINA")
print(f"  COORD 1: {cocinas['COORD 1']}  ·  COORD 2: {cocinas['COORD 2']}  "
      f"·  Gerencia: {cocinas['Gerencia']}")
print(f"  alguna cocina sin coordinadora: {'NO' if cocinas['Gerencia'] == 0 else 'SI'}")

print("\n[5] DESCANSO entre turnos (>= 12 h)")
violaciones = 0
for p in list(CONTRATO) + ["Gerencia"]:
    por_dia = {}
    for d, lista in turnos.items():
        for per, _, rs in lista:
            if per == p:
                por_dia[d] = rs
    for d in sorted(por_dia):
        if d + timedelta(days=1) not in por_dia:
            continue
        fin = max(b for _, b in por_dia[d])
        sig = min(a for a, _ in por_dia[d + timedelta(days=1)])
        if fin > 15:
            fin_abs = fin
            ini_abs = sig + 24
            if ini_abs - fin_abs < 12:
                violaciones += 1
                if violaciones <= 5:
                    print(f"    {p} {d}: sale {fin_abs:.2f} y entra {sig:.2f} -> "
                          f"{ini_abs - fin_abs:.2f} h")
print(f"  descansos por debajo de 12 h: {violaciones}")

# --- comparacion con la columna HORAS del propio libro
print("\n[6] CONTRASTE con la columna HORAS del libro")
tot_libro = 0.0
for hoja in wb.sheetnames:
    partes = hoja.split("-")
    if len(partes) != 2 or partes[1] not in MESES:
        continue
    ws = wb[hoja]
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=9):
        a = fila[0].value
        if isinstance(a, str) and a.strip() in list(CONTRATO) + ["Gerencia"]:
            v = fila[8].value
            if isinstance(v, (int, float)):
                tot_libro += float(v)
tot_texto = sum(horas.values())
print(f"  suma de la columna HORAS: {tot_libro:.2f}")
print(f"  suma recalculada del texto: {tot_texto:.2f}")
print(f"  diferencia: {abs(tot_libro - tot_texto):.2f} h")
