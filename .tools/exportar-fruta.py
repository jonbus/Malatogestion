"""Exporta la clasificación de fruta y verdura a datos para la aplicación."""
import json
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook("docs/fruta-verdura-familias.xlsx", data_only=True)
ws = wb["Clasificar"]

FAMILIAS = ["Fruta", "Verdura", "Embolsado", "Ensaladas Preparadas", "Legumbres", "Varios"]
productos = []
for r in range(5, ws.max_row + 1):
    ci = ws.cell(row=r, column=1).value
    prod = ws.cell(row=r, column=2).value
    fam = ws.cell(row=r, column=3).value
    if not prod:
        continue
    productos.append({
        "ci": "" if ci is None else str(ci).strip(),
        "nombre": str(prod).strip(),
        "familia": str(fam).strip() if fam else "",
    })

# --- problemas a revisar
raros = [p for p in productos if p["ci"] in ("96", "")]
dup = {k: v for k, v in Counter(p["ci"] for p in productos if len(p["ci"]) == 6).items() if v > 1}
con_ci_en_nombre = [p for p in productos if "(" in p["nombre"]]

datos = {
    "generado_de": "docs/fruta-verdura-familias.xlsx (completado por la propiedad)",
    "familias": FAMILIAS,
    "aviso_ci": "El CI es interno: solo gerencia y coordinadoras. Sirve para los pedidos.",
    "productos": productos,
}

with open("docs/datos/fruta-verdura.json", "w", encoding="utf-8") as f:
    json.dump(datos, f, ensure_ascii=False, indent=2)

print(f"  escrito docs/datos/fruta-verdura.json")
print(f"  productos: {len(productos)}")
c = Counter(p["familia"] for p in productos)
for fam in FAMILIAS:
    print(f"    {fam:<24} {c[fam]:>3}")
print()
print(f"  === CI que hay que revisar ({len(raros)}) ===")
for p in raros:
    print(f"    CI '{p['ci']}'  ·  {p['nombre']:<30} ({p['familia']})")
print()
print(f"  === CI repetido ({len(dup)}) ===")
for k in dup:
    print(f"    {k}:")
    for p in productos:
        if p["ci"] == k:
            print(f"       · {p['nombre']} ({p['familia']})")
print()
print(f"  === nombres con un número entre paréntesis ({len(con_ci_en_nombre)}) ===")
for p in con_ci_en_nombre:
    print(f"    {p['ci']:<8} {p['nombre']}  ({p['familia']})")
