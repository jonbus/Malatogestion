"""Mapa de la hoja: donde empieza cada seccion, cada semana y cada persona."""
import re
import sys
from openpyxl import load_workbook

RUTA = r"C:\Users\Usuario\Desktop\APP Malato Documentación\EjemploHorariosVacacionesVerano2026.xlsx"
wb = load_workbook(RUTA, data_only=False)
PERSONAS = {"OP 1", "OP 2", "OP 3", "COORD 1", "COORD 2", "YO"}

for hoja in wb.sheetnames:
    ws = wb[hoja]
    print("=" * 80)
    print(f"HOJA {hoja!r}  max_fila={ws.max_row} max_col={ws.max_column}")
    print("=" * 80)
    if ws.max_row <= 2:
        print("  (vacia)\n")
        continue
    for r in range(1, ws.max_row + 1):
        b = ws.cell(row=r, column=2).value
        a = ws.cell(row=r, column=1).value
        if b is None and a is None:
            continue
        if isinstance(b, str) and re.match(r"^\d{2}/\d{2}\s*-\s*\d{2}/\d{2}$", b.strip()):
            print(f"  r{r:>4}  SEMANA {b.strip()}")
        elif isinstance(b, str) and b.strip().upper() in PERSONAS:
            print(f"  r{r:>4}        - {b.strip()}")
        elif a is not None and str(a).strip():
            print(f"  r{r:>4}  [A] {str(a).strip()[:70]!r}")
        elif isinstance(b, str) and b.strip():
            print(f"  r{r:>4}  [B] {b.strip()[:70]!r}")
    print()
