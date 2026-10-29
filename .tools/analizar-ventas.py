"""Analiza el historico de ventas del asador: estructura, bloques y columnas extra."""
import sys
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook(r".tools\ventas-pollos.xlsx", data_only=True)

# --- 1) columnas extra en las hojas de 16 columnas
print("=" * 76)
print("COLUMNAS EXTRA (hojas de 16 columnas)")
print("=" * 76)
for hoja in ("2025", "2024", "2023"):
    ws = wb[hoja]
    print(f"\n  hoja {hoja}: max_col={ws.max_column}")
    for r in range(1, 20):
        o = ws.cell(row=r, column=15).value
        p = ws.cell(row=r, column=16).value
        q = ws.cell(row=r, column=17).value if ws.max_column >= 17 else None
        if any(x is not None and str(x).strip() for x in (o, p, q)):
            print(f"    r{r:>3}  O={o!r}  P={p!r}" + (f"  Q={q!r}" if q is not None else ""))

# --- 2) bloques por hoja
print("\n" + "=" * 76)
print("BLOQUES DE FINDE SEMANA POR AÑO")
print("=" * 76)
for hoja in ["2007", "2013", "2020", "2025", "2026"]:
    ws = wb[hoja]
    bloques = 0
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=2).value or "").strip().upper() == "M" and \
           str(ws.cell(row=r, column=1).value or "").strip() == "":
            bloques += 1
    print(f"  {hoja}: {bloques} bloques de fin de semana")

# --- 3) productos distintos que aparecen en la columna A
print("\n" + "=" * 76)
print("PRODUCTOS REGISTRADOS (columna A, fuera de cabeceras)")
print("=" * 76)
prods = defaultdict(int)
for hoja in wb.sheetnames:
    ws = wb[hoja]
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if isinstance(a, str) and a.strip() and a.strip().upper() not in ("VENTAS ASADOR",):
            t = a.strip()
            if len(t) <= 12 and not any(ch.isdigit() for ch in t):
                prods[t] += 1
for p, n in sorted(prods.items(), key=lambda x: -x[1]):
    print(f"  {n:>5}  {p}")

# --- 4) total de pollos por año, como muestra de la serie
print("\n" + "=" * 76)
print("TOTAL DE POLLOS POR AÑO (columna TOTAL, filas de POLLOS)")
print("=" * 76)
for hoja in wb.sheetnames:
    if not hoja.isdigit():
        continue
    ws = wb[hoja]
    tot = 0.0
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip().upper() == "POLLOS":
            v = ws.cell(row=r, column=14).value
            try:
                tot += float(v)
            except (TypeError, ValueError):
                pass
    print(f"  {hoja}: {tot:>8.0f} pollos")
