"""Vuelca celdas con su referencia exacta (COL+fila) para no suponer la estructura."""
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

RUTA = sys.argv[1]
HOJA = sys.argv[2]
DESDE = int(sys.argv[3])
HASTA = int(sys.argv[4])

wb = load_workbook(RUTA, data_only=False)
ws = wb[HOJA]
print(f"HOJA {HOJA!r} filas {DESDE}-{HASTA}")
print("(se listan solo las celdas con contenido, con su referencia)\n")

for r in range(DESDE, HASTA + 1):
    trozos = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if v is None or str(v).strip() == "":
            continue
        trozos.append(f"{get_column_letter(c)}{r}={v!r}")
    if trozos:
        print("  " + "  ".join(trozos))
